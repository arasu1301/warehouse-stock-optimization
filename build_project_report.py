from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("output/Warehouse_Stock_Optimization_Project_Report.docx")

NAVY = "17365D"
BLUE = "DCE6F1"
GRID = "D9E2F3"
GREY = "F3F6F9"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def border_cell(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border_cell(cell)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 9)
        shade(t.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, None, 8.5)
            if ridx % 2: shade(cells[i], GREY)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text); r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(value)

def code_block(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(0)
        shade_p = OxmlElement("w:shd")
        ppr = p._p.get_or_add_pPr(); ppr.append(shade_p)
        shade_p.set(qn("w:fill"), "F3F6F9")
        r = p.add_run(line)
        r.font.name = "Courier New"; r._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New"); r.font.size = Pt(7.5)

def page(doc): doc.add_page_break()

def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = "WAREHOUSE STOCK OPTIMIZATION SYSTEM"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Warehouse Stock Optimization System | Project Report | ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for r in footer.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("666666")

def build():
    OUT.parent.mkdir(exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)
    add_header_footer(sec)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"; styles["Title"].font.size = Pt(28); styles["Title"].font.bold = True
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Aptos Display"; styles[name].font.color.rgb = RGBColor(0,0,0)
    styles["Heading 1"].font.size = Pt(16); styles["Heading 2"].font.size = Pt(12)

    # Cover
    for _ in range(4): doc.add_paragraph()
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Warehouse Stock Optimization System")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run("Project Report"); r.bold = True; r.font.size = Pt(16)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run("A warehouse inventory, FEFO, alerting, reporting, and demand forecasting solution").italic = True
    for _ in range(6): doc.add_paragraph()
    cover = doc.add_table(rows=4, cols=2); cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_data = [("Prepared by", "Arasu P"), ("Programme", "B.Tech Information Technology"), ("Institution", "Loyola Institute of Technology"), ("Academic Year", "2025 - 2026")]
    for row, (a,b) in zip(cover.rows, cover_data):
        set_cell_text(row.cells[0], a, True, "FFFFFF", 10); shade(row.cells[0], NAVY)
        set_cell_text(row.cells[1], b, False, None, 10)
    doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Submitted as an academic and portfolio project").font.size=Pt(10)
    page(doc)

    heading(doc, "Declaration")
    para(doc, "I, Arasu P, declare that this report presents the design and implementation of the Warehouse Stock Optimization System developed as an academic and portfolio project. The work is prepared for demonstration of software engineering, data management, and applied analytics skills.")
    doc.add_paragraph("\n\nSignature: ________________________\nDate: _____________________________")
    heading(doc, "Acknowledgement")
    para(doc, "I thank the faculty, mentors, peers, and learning resources that supported my development of practical skills in Python, databases, web APIs, and applied artificial intelligence. Their guidance helped shape this project into a solution focused on a real operational problem.")
    page(doc)

    heading(doc, "Abstract")
    para(doc, "The Warehouse Stock Optimization System is a software solution for managing products, batch-wise inventory, warehouse locations, stock movements, expiry risk, and replenishment decisions. The system records incoming batches, tracks current quantities, dispatches orders using First Expiry First Out (FEFO), and preserves an audit trail for receipts, sales, and adjustments. It also imports historical sales data to estimate future demand, recommend reorder quantities, and identify unusual sales patterns.")
    para(doc, "The application uses a FastAPI backend, SQLite database, Streamlit dashboard, and explainable analytics methods. The included historical dataset contains 3,650 sales records across 10 product identifiers. The solution demonstrates an end-to-end workflow from product setup and stock receipt to dispatch, alerting, reporting, and demand-informed planning. It is suitable as a mini-project, interview portfolio artifact, and foundation for a more integrated warehouse platform.")
    heading(doc, "Keywords", 2); para(doc, "Warehouse management, inventory optimization, FEFO, batch tracking, demand forecasting, reorder recommendation, anomaly detection, FastAPI, SQLite, Streamlit.")
    page(doc)

    heading(doc, "Table of Contents")
    toc = [("1", "Introduction and Problem Statement"), ("2", "Objectives Scope and Requirements"), ("3", "System Architecture and Design"), ("4", "Database Design"), ("5", "Functional Modules"), ("6", "Algorithms and Analytics"), ("7", "Implementation Details"), ("8", "Testing and Results"), ("9", "User Guide"), ("10", "Limitations and Future Enhancements"), ("11", "Conclusion"), ("Appendix A", "API Reference")]
    table(doc, ["Section", "Content"], toc, [1.1, 5.7])
    page(doc)

    heading(doc, "1 Introduction and Problem Statement")
    para(doc, "Warehouses must maintain accurate stock records while handling multiple products, locations, batches, expiry dates, customer orders, and replenishment decisions. Manual registers and disconnected spreadsheets make it difficult to determine what is available, which batch should be dispatched, and when stock should be reordered. These gaps can cause stock-outs, excess stock, expired goods, and incomplete traceability.")
    para(doc, "This project addresses those challenges through a centralized warehouse system. Each product is identified by an SKU and optional barcode. Physical stock is maintained at batch level with its expiry date and storage location. Transactions are captured as movements so that operations can be reviewed later. The system additionally uses historical sales data to support informed decisions instead of relying only on manual judgement.")
    heading(doc, "Problem Statement", 2)
    para(doc, "How can a warehouse manage batch-level stock accurately, dispatch the correct batch based on expiry, identify stock risks early, and use previous sales patterns to support replenishment planning?")
    heading(doc, "2 Objectives Scope and Requirements")
    heading(doc, "Objectives", 2)
    bullets(doc, ["Maintain a single source of truth for products, batches, locations, quantities, and stock value.", "Dispatch suitable stock using FEFO so that earliest-expiring usable batches leave first.", "Provide movement traceability for receipt, sale, retrieval, placement, and adjustment events.", "Generate low-stock and near-expiry alerts to reduce operational risk.", "Use historical sales to forecast demand, recommend replenishment quantities, and highlight unusual sales values.", "Provide accessible REST API documentation and a lightweight operational dashboard."])
    heading(doc, "Scope", 2)
    table(doc, ["Included", "Outside current scope"], [["Product master data, barcodes, batch receipt, location placement, FEFO sale dispatch, adjustments, alerts, reports, CSV import, analytics", "Multi-warehouse transfers, supplier purchase orders, payment processing, hardware scanner drivers, enterprise identity integration, automated purchasing"]], [3.4, 3.4])
    page(doc)

    heading(doc, "3 System Architecture and Design")
    para(doc, "The solution follows a layered design. The user interacts through the FastAPI Swagger interface or the Streamlit dashboard. API routes validate requests and call service classes. Services enforce warehouse rules and store data through SQLite. The analytics layer reads sales history and returns explainable forecasts, reorder suggestions, and anomaly results.")
    table(doc, ["Layer", "Components", "Responsibility"], [["Presentation", "Swagger API documentation, Streamlit dashboard", "Product setup, stock receipt, dispatch, visibility of alerts and inventory"], ["API", "Products, batches, inventory, sales, placement, movements, alerts, reports, data import", "Request validation and HTTP responses"], ["Business services", "Stock, batch, sales, placement, import, alert services", "FEFO, quantity protection, audit logging, import rules"], ["Data and analytics", "SQLite, forecasting, reorder prediction, anomaly detection", "Persistent storage and decision support"]], [1.1, 2.6, 3.1])
    heading(doc, "Operational Flow", 2)
    table(doc, ["Step", "Flow"], [["1", "Create or import product master data"], ["2", "Receive batch with quantity, expiry date, and location"], ["3", "View inventory and address low-stock or expiry alerts"], ["4", "Dispatch customer sale through FEFO"], ["5", "Save movement and add confirmed sale to sales history"], ["6", "Use forecast and reorder recommendation for planning"]], [0.8, 6.0])
    page(doc)

    heading(doc, "4 Database Design")
    para(doc, "SQLite is used as the project database because it is lightweight, portable, and appropriate for an academic demonstration. The schema separates product master data, stock batches, movements, and sales history so that inventory state and transaction history remain traceable.")
    table(doc, ["Table", "Key fields", "Purpose"], [["products", "sku, name, barcode, reorder_level, unit_cost", "Stores product master data and reorder settings"], ["batches", "sku, batch_code, expiry_date, quantity, location", "Stores batch-level physical inventory"], ["movements", "sku, batch_code, quantity, movement_type, reference, created_at", "Provides a receipt, sale, retrieval, and adjustment audit trail"], ["sales_history", "sku, sale_date, quantity", "Stores historical and new confirmed sales for analytics"], ["users", "username, password_hash, role", "Reserved for user management and role expansion"]], [1.1, 2.9, 2.8])
    heading(doc, "Data Integrity Controls", 2)
    bullets(doc, ["SKU and barcode values are unique where supplied.", "Batch quantity cannot become negative.", "A product must exist before its batch can be received.", "A sale is rejected when unexpired available quantity is insufficient.", "The database enables foreign-key enforcement for related data.", "CSV import validates required columns, dates, and numeric quantities."])
    page(doc)

    heading(doc, "5 Functional Modules")
    table(doc, ["Module", "Functions", "Business Value"], [["Product management", "Create, view, update SKU, barcode, reorder level, unit cost", "Maintains consistent master data"], ["Batch and inventory", "Receive batches, record location, view stock value and expiry sequence", "Provides accurate batch-level inventory"], ["Placement and adjustment", "Move a batch to a location, add or remove stock with a reason", "Supports warehouse corrections and traceability"], ["Sales and FEFO", "Dispatch sales from earliest-expiring usable batches", "Reduces expiry-related waste"], ["Movements", "Record receipt, sale, retrieval, adjustment references", "Creates an operational audit trail"], ["Alerts and reports", "Identify low stock and upcoming expiry; summarize value", "Helps staff prioritize actions"], ["Data import and AI", "Import CSV, forecast demand, suggest reorder quantity, detect anomalies", "Supports data-informed inventory planning"]], [1.35, 3.0, 2.45])
    heading(doc, "FEFO Business Rule", 2)
    para(doc, "The system filters batches for the requested SKU where quantity is positive and expiry date is today or later. It sorts those batches by expiry date and then by batch identifier. The requested quantity is taken from the first batch and, only if necessary, continues to later batches. Every selected quantity creates a movement record. This approach is especially useful for groceries, healthcare supplies, food, and any inventory with expiry constraints.")
    page(doc)

    heading(doc, "6 Algorithms and Analytics")
    heading(doc, "Demand Forecasting", 2)
    para(doc, "The forecasting module applies a least-squares linear trend to chronological sales quantities. It returns a daily forecast and estimated demand for a selected future period. This method is transparent and suitable for a baseline where the user needs to understand why a recommendation was generated.")
    heading(doc, "Reorder Recommendation", 2)
    para(doc, "The target stock quantity is calculated from forecast demand across the configured lead time and safety days. Recommended order quantity is the positive difference between target stock and current available warehouse stock. This prevents negative recommendations and makes assumptions visible to the user.")
    heading(doc, "Anomaly Detection", 2)
    para(doc, "The anomaly module uses the median and median absolute deviation (MAD), which are robust against isolated extreme values. When typical sales are constant, any different value is treated as a noteworthy anomaly. The result highlights records that may require review for promotions, data-entry mistakes, unusual demand, or supply changes.")
    heading(doc, "Representative Data Result", 2)
    table(doc, ["Metric", "Observed value"], [["Imported sales records", "3,650"], ["Distinct product identifiers", "10"], ["P001 sales-history records", "365"], ["P001 estimated daily demand", "33.33 units"], ["P001 estimated 7-day demand", "233.31 units"], ["Forecasting method", "Linear trend"]], [3.2, 3.6])
    para(doc, "These values demonstrate that the data pipeline and analytics endpoints are connected. Forecast quality should be evaluated periodically against future actual sales before any automatic purchasing decision is made.")
    page(doc)

    heading(doc, "7 Implementation Details")
    table(doc, ["Technology", "Use in project"], [["Python", "Core application language and analytics implementation"], ["FastAPI", "REST API, validation, interactive Swagger documentation"], ["SQLite", "Embedded relational database for products, batches, movements, and sales history"], ["Pydantic", "Request validation and safe API schemas"], ["Streamlit", "Operational dashboard for inventory, receipt, and sales dispatch"], ["CSV", "Historical sales data import format"], ["Pytest", "Automated unit tests for core workflows"]], [2.0, 4.8])
    heading(doc, "Key API Endpoints", 2)
    table(doc, ["Method", "Endpoint", "Purpose"], [["POST", "/products", "Create product"], ["PUT", "/products/{sku}", "Update product master data"], ["POST", "/batches", "Receive a stock batch"], ["POST", "/placement", "Change a batch location"], ["POST", "/inventory/adjust", "Apply controlled stock adjustment"], ["POST", "/sales/dispatch", "Dispatch FEFO sale and update sales history"], ["GET", "/alerts", "View low-stock and expiry alerts"], ["GET", "/reports/inventory", "View units and inventory value"], ["POST", "/data/import-historical-sales", "Import historical sales CSV"], ["GET", "/ai/forecast/{sku}", "Get forecast for a product"]], [0.8, 2.4, 3.6])
    page(doc)

    heading(doc, "8 Testing and Results")
    para(doc, "The project includes automated tests for product validation, batch validation, FEFO selection, forecasting output, anomaly behavior, data import protection, stock adjustment safety, and sale recording. Tests are run with the command python -m pytest -q from the project virtual environment.")
    table(doc, ["Test area", "Expected result"], [["Product model", "Accepts valid SKU and name; rejects invalid required values"], ["Batch model", "Rejects past expiry dates and non-positive quantity"], ["FEFO retrieval", "Uses earliest expiry batch before later batch"], ["Sales dispatch", "Reduces batch quantity, creates SALE movement, writes sales history"], ["Stock adjustment", "Rejects any adjustment that would make quantity negative"], ["CSV import", "Imports valid source once and skips a duplicate repeat"], ["Anomaly detection", "Flags an exceptional value against a constant baseline"]], [2.3, 4.5])
    heading(doc, "Manual Verification Checklist", 2)
    bullets(doc, ["Create or update product P001 and configure a reorder level.", "Receive at least two batches with different future expiry dates.", "Dispatch a sale larger than the first batch only when a split should occur.", "Confirm that the earliest expiry batch is consumed first in the response and movement log.", "Check alerts, inventory report, and AI endpoints in Swagger documentation.", "Open the Streamlit dashboard and verify receipt and sale forms return expected API responses."])
    page(doc)

    heading(doc, "9 User Guide")
    heading(doc, "Starting the Application", 2)
    para(doc, "Activate the project virtual environment, install dependencies, and start the server. Open http://127.0.0.1:8000/docs for interactive API testing. In a second terminal, run streamlit run frontend/dashboard.py to open the dashboard.")
    heading(doc, "Using Historical Sales Data", 2)
    para(doc, "The file data/historical_sales.csv is imported using POST /data/import-historical-sales. The CSV requires date, product_id, product_name, quantity_sold, and unit_price columns. The imported product identifier becomes the SKU used by forecast and reorder endpoints. For example, use P001 in GET /ai/forecast/P001.")
    heading(doc, "Using the System Without a Live Sales Integration", 2)
    para(doc, "Live data is not required for a project demonstration. Existing historical data establishes the initial forecast. New sales can be recorded one by one using POST /sales/dispatch or from the dashboard; each confirmed sale is added to the analytics history. A business can also export sales from Excel, a point-of-sale system, an ERP, or an e-commerce platform on a weekly or monthly basis and replace the CSV. The import endpoint supports a controlled replace_existing option when a complete refreshed history is available.")
    page(doc)

    heading(doc, "10 Limitations and Future Enhancements")
    table(doc, ["Current limitation", "Recommended enhancement"], [["Single embedded database", "Use PostgreSQL or MySQL for concurrent multi-user deployment"], ["Manual/periodic sales import", "Integrate POS, ERP, billing, or e-commerce webhooks"], ["Baseline linear forecast", "Evaluate seasonal forecasting models with back-testing and accuracy metrics"], ["No purchase-order workflow", "Add suppliers, purchase orders, approvals, and receiving reconciliation"], ["Basic dashboard", "Add role-based views, charts, exports, and mobile-friendly scanning"], ["Software barcode lookup", "Connect a USB/mobile scanner and label-printing workflow"], ["User table reserved", "Implement authentication, authorization, and audit permissions"]], [2.6, 4.2])
    para(doc, "These enhancements are intentionally separated from the current scope. The implemented project provides a strong working foundation while leaving clear paths for enterprise-scale deployment.")
    heading(doc, "11 Conclusion")
    para(doc, "The Warehouse Stock Optimization System demonstrates a complete operational workflow for product and batch management, FEFO dispatch, movement traceability, alerts, reporting, and analytics-based planning. It converts historical sales records into meaningful demand and reorder insights while retaining practical warehouse controls such as expiry tracking and negative-stock protection. The project is suitable for academic evaluation and job interviews because it combines backend development, database design, user-facing tools, business logic, data processing, and explainable AI features in one coherent application.")
    page(doc)

    heading(doc, "Appendix A API Reference")
    table(doc, ["Area", "Endpoint", "Typical use"], [["System", "GET / and GET /health", "Confirm that the service is running"], ["Products", "POST /products, GET /products, PUT /products/{sku}", "Manage master data and reorder settings"], ["Inventory", "GET /inventory, POST /inventory/retrieve, POST /inventory/adjust", "View, retrieve, or correct stock"], ["Batches", "POST /batches, GET /batches", "Receive and inspect stock batches"], ["Sales", "POST /sales/dispatch", "Dispatch FEFO order and save confirmed sale"], ["Placement", "POST /placement", "Move a batch to another warehouse location"], ["Monitoring", "GET /movements, GET /alerts, GET /reports/inventory", "Review traceability, risks, and value"], ["Analytics", "GET /ai/forecast/{sku}, /ai/reorder/{sku}, /ai/anomalies/{sku}", "Plan stock using historical sales"], ["Data", "POST /data/import-historical-sales", "Load initial or replacement historical sales data"]], [1.2, 3.4, 2.2])
    heading(doc, "Appendix B References")
    bullets(doc, ["FastAPI documentation: https://fastapi.tiangolo.com/", "Streamlit documentation: https://docs.streamlit.io/", "SQLite documentation: https://www.sqlite.org/docs.html", "Pydantic documentation: https://docs.pydantic.dev/", "Project source code, database schema, and historical sales dataset in the Warehouse Stock Optimization workspace."])
    page(doc)

    heading(doc, "Appendix C Key Code Implementation")
    para(doc, "This appendix presents selected implementation excerpts that demonstrate the principal warehouse and analytics logic. The complete executable source code is maintained in the project workspace. The excerpts are intentionally focused on business rules rather than repeating framework boilerplate.")
    heading(doc, "C1 FEFO Stock Dispatch", 2)
    para(doc, "The stock service selects unexpired positive-quantity batches in ascending expiry order. It rejects insufficient stock and logs every selected batch as a movement. This is the core rule that prevents later-expiring stock from being dispatched before earlier-expiring stock.")
    code_block(doc, '''rows = c.execute(
    "SELECT * FROM batches WHERE sku=? AND quantity>0 AND expiry_date>=? "
    "ORDER BY expiry_date, id", (request.sku, str(date.today()))
).fetchall()
if sum(r["quantity"] for r in rows) < remaining:
    raise HTTPException(400, "Insufficient unexpired stock")
for row in rows:
    take = min(remaining, row["quantity"])
    c.execute("UPDATE batches SET quantity=quantity-? WHERE id=?", (take, row["id"]))
    c.execute("INSERT INTO movements(...) VALUES(...)" )
    remaining -= take''')
    heading(doc, "C2 Sale Dispatch and Analytics History", 2)
    para(doc, "A customer order is dispatched through the sales service, which reuses FEFO and stores the confirmed quantity in sales history in the same database transaction. This causes future forecasts to include newly recorded sales.")
    code_block(doc, '''class SalesService(StockService):
    def dispatch_sale(self, request):
        return self._retrieve_fefo(
            request,
            movement_type="SALE",
            record_sale=True,
        )

if record_sale:
    c.execute(
        "INSERT INTO sales_history(sku,sale_date,quantity) VALUES(?,?,?)",
        (request.sku, str(date.today()), request.quantity),
    )''')
    heading(doc, "C3 Historical CSV Import", 2)
    para(doc, "The importer validates required CSV columns, validates dates and numeric quantities, creates missing product master records, and stores each sales row for the analytics module. It skips duplicate full imports unless the replace-existing option is explicitly supplied.")
    code_block(doc, '''reader = csv.DictReader(source)
if not REQUIRED_COLUMNS.issubset(reader.fieldnames):
    raise HTTPException(400, "CSV columns are invalid")
date.fromisoformat(row["date"].strip())
quantity = int(float(row["quantity_sold"]))
connection.execute(
    "INSERT OR IGNORE INTO products(sku,name,reorder_level,unit_cost) VALUES(?,?,?,?)",
    (sku, row["product_name"].strip(), 0, unit_cost),
)
connection.execute(
    "INSERT INTO sales_history(sku,sale_date,quantity) VALUES(?,?,?)",
    (sku, row["date"].strip(), quantity),
)''')
    heading(doc, "C4 Forecast and Reorder Calculation", 2)
    para(doc, "The baseline forecast computes a linear trend from chronological quantity values. The reorder module converts predicted demand into a target stock level using lead time and safety days, then recommends only the positive shortage.")
    code_block(doc, '''slope = sum((x-xbar)*(y-ybar) for x, y in zip(xs, values)) / divisor
prediction = max(0, ybar + slope * len(values))

target = round(model["daily_forecast"] * (lead_time_days + safety_days))
recommended_order = max(0, target - current_stock)''')
    heading(doc, "C5 Robust Anomaly Detection", 2)
    para(doc, "Median absolute deviation is used to make anomaly detection less sensitive to a small number of extreme values. The fallback logic also identifies an unusual value when a normally constant series changes suddenly.")
    code_block(doc, '''centre = median(values)
mad = median([abs(x-centre) for x in values])
if mad == 0:
    return [row for row in history if float(row["quantity"]) != centre]
score = 0.6745 * (float(row["quantity"]) - centre) / mad
if abs(score) > threshold:
    anomalies.append(row)''')
    page(doc)

    heading(doc, "Appendix D Complete Source Code Inventory")
    para(doc, "The following inventory lists the project source files and their role. This provides a complete map of the codebase while keeping the report focused on essential implementation excerpts. All files are contained in the Warehouse Stock Optimization project folder.")
    table(doc, ["Folder", "Files", "Purpose"], [
        ("Root", "run.py, requirements.txt, README.md, PROJECT_PLAN.md", "Application launch, dependencies, setup instructions, and project plan"),
        ("app", "main.py, config.py, __init__.py", "FastAPI application startup, route registration, database initialization, configuration"),
        ("app/models", "product.py, batch.py, stock.py, movement.py, user.py", "Pydantic request schemas and data models"),
        ("app/database", "database.py, tables.py", "SQLite connection handling, schema, indexes, foreign-key enforcement"),
        ("app/services", "stock_service.py, sales_service.py, purchase_service.py, placement_service.py, retrieval_service.py, batch_service.py, alert_service.py, data_import_service.py", "Warehouse business rules, FEFO, movements, alerts, data import"),
        ("app/api", "products.py, batches.py, inventory.py, sales.py, placement.py, movements.py, alerts.py, reports.py, data.py, auth.py", "REST endpoints grouped by functional area"),
        ("app/ai", "demand_forecasting.py, reorder_prediction.py, anomaly_detection.py", "Forecast, stock recommendation, and unusual-sales analysis"),
        ("app/barcode", "scanner.py", "Barcode-based product lookup"),
        ("app/utils", "validation.py, notifications.py", "Shared validation and alert-formatting helpers"),
        ("frontend", "dashboard.py and pages inventory, products, batches, movements, alerts, reports, ai_predictions", "Streamlit user interface and page structure"),
        ("data", "historical_sales.csv, warehouse.db", "Historical input dataset and local application database"),
        ("tests", "test_products.py, test_batches.py, test_inventory.py, test_retrieval.py, test_ai.py, test_operations.py, test_data_import.py", "Automated checks for models, FEFO, sales, import, adjustment, and analytics")
    ], [1.05, 3.4, 2.55])
    heading(doc, "Coding Standards Applied", 2)
    bullets(doc, ["Layered separation between API routes, business services, models, and database code.", "Pydantic validation for input shape, length, quantity, and expiry constraints.", "Parameterized SQL queries to avoid concatenating user values into statements.", "Database transactions for stock-changing operations and sale-history recording.", "Explicit errors for duplicate master data, insufficient stock, missing batches, invalid CSV content, and negative inventory.", "Automated tests for the high-risk rules: FEFO order, stock safety, import idempotency, and anomaly behavior."])
    doc.core_properties.title = "Warehouse Stock Optimization System Project Report"
    doc.core_properties.author = "Arasu P"
    doc.core_properties.subject = "Warehouse Stock Optimization System"
    doc.save(OUT)

if __name__ == "__main__": build()
